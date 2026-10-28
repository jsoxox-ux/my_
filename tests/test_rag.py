from unittest.mock import Mock, patch
import pytest
from docx import Document
from rag import Chunk, Index, load_folder, answer, split_text


def test_korean_retrieval_and_unknown(tmp_path):
    (tmp_path / '학칙.txt').write_text('휴학 신청은 학기 시작 전까지 교무처에 제출합니다. 휴학 신청 서류는 신청서입니다.', encoding='utf-8')
    (tmp_path / '장학.md').write_text('장학금 지급 대상은 성적 우수 학생입니다.', encoding='utf-8')
    chunks, warnings = load_folder(tmp_path)
    assert not warnings
    index = Index(chunks)
    assert index.search('휴학 신청 서류')[0][0].source == '학칙.txt'
    assert index.search('zzzzzz') == []


def test_docx_tables_and_bad_pdf(tmp_path):
    doc = Document()
    doc.add_paragraph('복학 규정')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '복학 신청서는 교무처 제출'
    doc.save(tmp_path / '복학.docx')
    (tmp_path / '손상.pdf').write_bytes(b'not a pdf')
    chunks, warnings = load_folder(tmp_path)
    assert '교무처 제출' in chunks[0].text
    assert len(warnings) == 1


def test_overlap_keeps_document():
    original = ''.join(str(i % 10) for i in range(2600))
    parts = list(split_text(original))
    assert parts[0] + ''.join(p[200:] for p in parts[1:]) == original


@pytest.mark.parametrize('provider', ['Claude', 'Gemini'])
def test_provider_request(provider):
    data = {'content': [{'type': 'text', 'text': '답변 [1]'}]} if provider == 'Claude' else {'candidates': [{'content': {'parts': [{'text': '답변 [1]'}]}}]}
    response = Mock(status_code=200)
    response.json.return_value = data
    with patch('rag.requests.post', return_value=response) as post:
        assert answer(provider, 'test-model', 'test-secret', '휴학?', [(Chunk('학칙.txt', '본문', '휴학 규정'), .8)]) == '답변 [1]'
        assert 'test-secret' not in str(post.call_args.kwargs['json'])
        assert '학칙.txt' in str(post.call_args.kwargs['json'])


def test_api_error_hides_response_and_key():
    with patch('rag.requests.post', return_value=Mock(status_code=401, text='sensitive')):
        with pytest.raises(RuntimeError, match='API 키') as error:
            answer('Claude', 'test', 'secret', '질문', [])
        assert 'sensitive' not in str(error.value)


def test_streamlit_initial_screen():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file('../app.py').run()
    assert not app.exception
    assert app.chat_input[0].disabled
    assert app.selectbox[0].value == 'Claude'
    app.selectbox[0].select('Gemini').run()
    assert not app.exception
    assert app.text_input[0].value == 'gemini-2.5-flash'


def test_chat_with_document(tmp_path):
    from streamlit.testing.v1 import AppTest
    (tmp_path / '규정.txt').write_text('휴학 신청은 교무처에 제출합니다.', encoding='utf-8')
    app = AppTest.from_file('../app.py').run()
    app.text_input[2].set_value(str(tmp_path))
    app.button[0].click().run()
    assert not app.exception
    assert not app.chat_input[0].disabled
    app.text_input[1].set_value('test-key').run()
    with patch('rag.answer', return_value='교무처에 제출합니다. [1]'):
        app.chat_input[0].set_value('휴학 신청은 어디에?').run()
    assert not app.exception
    assert len(app.chat_message) == 2
    assert '교무처' in app.session_state.messages[-1]['content']
    assert app.session_state.messages[-1]['sources'][0][0].source == '규정.txt'
